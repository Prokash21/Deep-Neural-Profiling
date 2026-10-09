# Build manuscript.docx and manuscript.html from manuscript_content.py.
# PDF: print manuscript.html with Chrome (the page numbers its own lines with JavaScript before printing):
#   "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome" --headless --no-pdf-header-footer \
#       --virtual-time-budget=10000 --print-to-pdf="$PWD/manuscript.pdf" "file://$PWD/manuscript.html"
# Requires: python-docx (pip install python-docx)
import csv
import html
import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

import manuscript_content as M

HERE = os.path.dirname(os.path.abspath(__file__))


def table_data(key):
    cap, header, rows = M.TABLES[key]
    if isinstance(rows, str) and rows.startswith("CSV:"):
        with open(os.path.join(HERE, rows[4:]), encoding="utf8") as f:
            rows = [[r["GEO_Accession"], r["Samples_in_series"], r["Platform"], r["Study_Type"]] for r in csv.DictReader(f)]
    return cap, header, rows


def split_caption(cap):
    label, text = cap.split(". ", 1)
    return label + ".", text


# ---------------- DOCX ----------------
def add_line_numbers(section):
    """Continuous line numbering in the left margin (Word: Layout > Line Numbers)."""
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1")
    ln.set(qn("w:restart"), "continuous")
    ln.set(qn("w:distance"), "340")
    pg_mar = section._sectPr.find(qn("w:pgMar"))
    pg_mar.addnext(ln)


def no_line_numbers(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    ppr.append(OxmlElement("w:suppressLineNumbers"))


def add_caption(doc, cap, size=10):
    label, text = split_caption(cap)
    p = doc.add_paragraph()
    p.add_run(label + " ").bold = True
    p.add_run(text).font.size = Pt(size)
    return p


def add_figure(doc, key):
    img, cap = M.FIGURES[key]
    doc.add_picture(os.path.join(HERE, img), width=Cm(16))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    no_line_numbers(doc.paragraphs[-1])
    add_caption(doc, cap)


def build_docx(path):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(11)
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(2.2)
        sec.top_margin = sec.bottom_margin = Cm(2.2)
        add_line_numbers(sec)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(M.TITLE)
    r.bold = True
    r.font.size = Pt(15)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for i, (name, sup) in enumerate(M.AUTHORS):
        p.add_run(name)
        p.add_run(sup).font.superscript = True
        if i < len(M.AUTHORS) - 1:
            p.add_run(", ")
    for line in (M.AFFIL, M.CORRESP, M.VERSION):
        q = doc.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if line == M.AFFIL:
            q.add_run(line[0]).font.superscript = True
            q.add_run(line[1:].lstrip())
        else:
            q.add_run(line)

    doc.add_heading("Abstract", level=1)
    for label, text in M.ABSTRACT:
        p = doc.add_paragraph()
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.add_run(label + ": ").bold = True
        p.add_run(text)
    p = doc.add_paragraph()
    p.add_run("Keywords: ").bold = True
    p.add_run(M.KEYWORDS)

    for heading, blocks in M.SECTIONS:
        doc.add_heading(heading, level=1)
        for kind, val in blocks:
            if kind == "p":
                p = doc.add_paragraph(val)
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            elif kind == "h3":
                doc.add_heading(val, level=2)
            elif kind == "fig":
                add_figure(doc, val)
            elif kind == "table":
                cap, header, rows = table_data(val)
                add_caption(doc, cap)
                tb = doc.add_table(rows=1, cols=len(header))
                tb.style = "Light Grid Accent 1"
                for c, h in zip(tb.rows[0].cells, header):
                    c.text = h
                    c.paragraphs[0].runs[0].bold = True
                for r in rows:
                    for c, v in zip(tb.add_row().cells, r):
                        c.text = v
                for row_ in tb.rows:
                    for c in row_.cells:
                        no_line_numbers(c.paragraphs[0])
                        for run in c.paragraphs[0].runs:
                            run.font.size = Pt(8.5)
                doc.add_paragraph()

    doc.add_heading("Data and code availability", level=1)
    doc.add_paragraph(M.DATA_AVAIL)
    doc.add_heading("Supplementary material", level=1)
    add_figure(doc, "figS1")
    doc.add_paragraph("Supplementary Tables S1–S10 (supplementary/Supplementary_Tables_S1-S10.xlsx):")
    for code, desc in M.SUPP_TABLES:
        doc.add_paragraph(f"Table {code}. {desc}", style="List Bullet")
    for title, text in (("Author contributions", M.CONTRIB), ("Acknowledgements", M.ACK),
                        ("Competing interests", M.COMPETING)):
        doc.add_heading(title, level=1)
        doc.add_paragraph(text)
    doc.add_heading("References", level=1)
    for ref in M.REFERENCES:
        p = doc.add_paragraph(ref)
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.first_line_indent = Cm(-0.6)
        for run in p.runs:
            run.font.size = Pt(10)
    doc.save(path)


# ---------------- HTML ----------------
CSS = """
@page { size: A4; margin: 20mm 16mm 20mm 14mm; }
body { font-family: 'Times New Roman', Times, serif; font-size: 11pt; line-height: 1.45; color: #111; width: 164mm; margin: 0 auto; padding-left: 12mm; }
h1.title { font-size: 17pt; text-align: center; margin: 0 0 8pt; line-height: 1.25; }
.meta { text-align: center; font-size: 10.5pt; margin: 2pt 0; }
h2 { font-size: 13pt; margin: 16pt 0 6pt; }
h3 { font-size: 11.5pt; margin: 12pt 0 4pt; }
p { text-align: justify; margin: 0 0 7pt; }
figure { margin: 12pt 0; page-break-inside: avoid; }
figure img { width: 100%; }
figcaption { font-size: 9.5pt; text-align: justify; }
.refs p { font-size: 9.5pt; padding-left: 1.5em; text-indent: -1.5em; margin-bottom: 3pt; text-align: left; }
ul { font-size: 10pt; }
.tbl { margin: 10pt 0; page-break-inside: avoid; }
.tcap { font-size: 9.5pt; margin-bottom: 3pt; }
table { border-collapse: collapse; width: 100%; font-size: 8.5pt; }
th, td { border-top: 1px solid #bbb; border-bottom: 1px solid #bbb; padding: 2pt 4pt; text-align: left; }
th { background: #eee; }
/* line numbers: a zero-width marker before the first word of every line */
.ln { position: relative; }
.ln::before { content: attr(data-n); position: absolute; left: calc(-1 * var(--dx) - 11mm); width: 9mm;
              text-align: right; font-size: 8pt; color: #777; font-weight: normal; font-style: normal; text-indent: 0; }
"""

# Numbers every text line of the body (abstract onwards, excluding table cells). Markers are inline, so they move
# with the text when the page is paginated for printing.
LINE_NUMBER_JS = r"""
<script>
(function () {
  const blocks = document.querySelectorAll('.numbered p, .numbered h2, .numbered h3, .numbered figcaption, .numbered li');
  const bs = getComputedStyle(document.body);
  const bodyLeft = document.body.getBoundingClientRect().left + parseFloat(bs.paddingLeft);
  const words = [];
  blocks.forEach(b => {
    const walker = document.createTreeWalker(b, NodeFilter.SHOW_TEXT);
    const nodes = []; while (walker.nextNode()) nodes.push(walker.currentNode);
    nodes.forEach(n => {
      const parts = n.textContent.split(/(\s+)/);
      const frag = document.createDocumentFragment();
      parts.forEach(t => {
        if (!t) return;
        if (/^\s+$/.test(t)) { frag.appendChild(document.createTextNode(t)); return; }
        const s = document.createElement('span'); s.textContent = t; frag.appendChild(s); words.push(s);
      });
      n.parentNode.replaceChild(frag, n);
    });
  });
  let lastTop = null, n = 0;
  words.forEach(w => {
    const r = w.getBoundingClientRect();
    if (lastTop === null || Math.abs(r.top - lastTop) > 3) {
      n += 1; lastTop = r.top;
      const m = document.createElement('span');
      m.className = 'ln'; m.dataset.n = n;
      m.style.setProperty('--dx', (r.left - bodyLeft) + 'px');
      w.parentNode.insertBefore(m, w);
    }
  });
})();
</script>
"""


def esc(s):
    s = html.escape(s)
    return re.sub(r"(https?://[^\s<]+)", r'<a href="\1">\1</a>', s)


def fig_html(key):
    img, cap = M.FIGURES[key]
    label, text = split_caption(cap)
    return f'<figure><img src="{img}"><figcaption><b>{esc(label)}</b> {esc(text)}</figcaption></figure>'


def build_html(path):
    out = [f"<!doctype html><html><head><meta charset='utf-8'><title>{esc(M.TITLE)}</title><style>{CSS}</style></head><body>"]
    out.append(f'<h1 class="title">{esc(M.TITLE)}</h1>')
    authors = ", ".join(f"{esc(n)}<sup>{esc(s)}</sup>" for n, s in M.AUTHORS)
    out.append(f'<p class="meta">{authors}</p>')
    out.append(f'<p class="meta"><sup>{esc(M.AFFIL[0])}</sup>{esc(M.AFFIL[1:].lstrip())}</p>')
    for line in (M.CORRESP, M.VERSION):
        out.append(f'<p class="meta">{esc(line)}</p>')
    out.append('<div class="numbered"><h2>Abstract</h2>')
    for label, text in M.ABSTRACT:
        out.append(f"<p><b>{esc(label)}:</b> {esc(text)}</p>")
    out.append(f"<p><b>Keywords:</b> {esc(M.KEYWORDS)}</p>")
    for heading, blocks in M.SECTIONS:
        out.append(f"<h2>{esc(heading)}</h2>")
        for kind, val in blocks:
            if kind == "p":
                out.append(f"<p>{esc(val)}</p>")
            elif kind == "h3":
                out.append(f"<h3>{esc(val)}</h3>")
            elif kind == "fig":
                out.append(fig_html(val))
            elif kind == "table":
                cap, header, rows = table_data(val)
                label, text = split_caption(cap)
                out.append(f'<div class="tbl"><p class="tcap"><b>{esc(label)}</b> {esc(text)}</p><table><tr>'
                           + "".join(f"<th>{esc(h)}</th>" for h in header) + "</tr>"
                           + "".join("<tr>" + "".join(f"<td>{esc(v)}</td>" for v in r) + "</tr>" for r in rows)
                           + "</table></div>")
    out.append(f"<h2>Data and code availability</h2><p>{esc(M.DATA_AVAIL)}</p>")
    out.append("<h2>Supplementary material</h2>" + fig_html("figS1"))
    out.append("<p>Supplementary Tables S1–S10 (supplementary/Supplementary_Tables_S1-S10.xlsx):</p><ul>")
    out += [f"<li><b>Table {c}.</b> {esc(d)}</li>" for c, d in M.SUPP_TABLES]
    out.append("</ul>")
    for title, text in (("Author contributions", M.CONTRIB), ("Acknowledgements", M.ACK),
                        ("Competing interests", M.COMPETING)):
        out.append(f"<h2>{title}</h2><p>{esc(text)}</p>")
    out.append('<h2>References</h2><div class="refs">' + "".join(f"<p>{esc(r)}</p>" for r in M.REFERENCES) + "</div>")
    out.append("</div>" + LINE_NUMBER_JS + "</body></html>")
    with open(path, "w", encoding="utf8") as f:
        f.write("\n".join(out))


if __name__ == "__main__":
    build_docx(os.path.join(HERE, "manuscript.docx"))
    build_html(os.path.join(HERE, "manuscript.html"))
    print("built manuscript.docx and manuscript.html")
